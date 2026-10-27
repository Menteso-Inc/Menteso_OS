import openpyxl
from agents.pct_agent import ai_verifier as ai, agent, contact_policy as policy, signer_evidence as signer
from agents.pct_agent.test_contact_ownership import contact
from agents.pct_agent.test_local_evidence import line


def block(owner, email, number, country='Germany', role='agent'):
    common = dict(entity=owner, block_id=f'contact-{number}',
                  section=f'IV-{number} Agent' if role == 'agent' else 'II Applicant', role=role)
    return [contact(kind, value, **common) for kind, value in [
        ('name', owner), ('email', email), ('country', country),
        ('phone', '+49 123456789' if number == 1 else '+41 987654321')]]


def signature(value, kind='signer_entity', **changes):
    return contact(kind, value, entity=value, role='other', section='Signature of applicant, agent or common representative',
                   block_id='signature-p2-1', **changes)


def select(*fields):
    return ai.select_contacts(ai.VisualExtraction(fields=list(fields), reviewed_pages=[1, 2], uncertainties=[]), ai.load_rules(), [1, 2])


def pair():
    return [*block('First Counsel LLP', 'person@first.example', 1),
            *block('Signing Counsel SA', 'filing@signing.example', 2, 'Switzerland')]


def test_signing_organisation_selects_whole_additional_contact_block():
    result = select(*pair(), signature('Signing Counsel SA'))
    assert result['emails'] == ['filing@signing.example']
    assert result['country'] == 'CH' and result['phones'] == ['+41 987654321']
    assert result['contact_owner'] == result['agent_name'] == 'Signing Counsel SA'
    assert result['additional_emails'][0]['email'] == 'person@first.example'
    assert result['additional_emails'][0]['owner'] == 'First Counsel LLP'
    assert result['signer_selection']['status'] == 'confirmed'


def test_signer_person_matches_only_complete_name_not_initial_or_domain():
    fields = [*block('First Counsel LLP', 'person@first.example', 1),
              *block('GREEN, Alex', 'alex@second.example', 2)]
    assert select(*fields, signature('Alex GREEN', kind='signer_name'))['emails'] == ['alex@second.example']
    assert select(*fields, signature('Alex', kind='signer_name'))['emails'] == ['person@first.example']


def test_unclear_or_unmatched_signature_retains_previous_selection():
    for observation in [signature('Signing Counsel SA', legible=False), signature('Unlisted Firm'),
                        signature('Signing Counsel SA', evidence='Unreadable'), signature('Signing Counsel SA', page=3)]:
        assert select(*pair(), observation)['emails'] == ['person@first.example']


def test_multiple_conflicting_signers_do_not_arbitrarily_choose_one():
    second = signature('Signing Counsel SA').model_copy(update={'block_id': 'signature-p2-2'})
    result = select(*pair(), signature('First Counsel LLP'), second)
    assert result['emails'] == ['person@first.example']
    assert result['signer_selection']['status'] == 'unresolved'


def test_additional_emails_have_owners_exclude_office_inventor_and_main_duplicates():
    result = select(*pair(), *block('Applicant Ltd', 'applicant@example.org', 3, role='applicant'),
                    contact('email', 'inventor@example.org', role='inventor'),
                    contact('email', 'office@wipo.int', role='patent_office'),
                    signature('Signing Counsel SA'))
    assert result['emails'] == ['applicant@example.org']
    assert {item['email'] for item in result['additional_emails']} == {'person@first.example', 'filing@signing.example'}
    assert {item['role'] for item in result['additional_emails']} == {'agent'}
    policy.withhold_email(result, 'applicant@example.org')
    assert policy.result_emails(result) == ['person@first.example', 'filing@signing.example']


def test_signing_agent_does_not_replace_preferred_applicant_email():
    result = select(*block('Applicant Ltd', 'applicant@example.org', 3, role='applicant'),
                    *block('Signing Counsel SA', 'filing@signing.example', 2), signature('Signing Counsel SA'))
    assert result['emails'] == ['applicant@example.org']
    assert result['signer_selection']['status'] == 'retained_role_policy'


def test_main_email_appearing_for_multiple_owners_is_not_duplicated_as_additional():
    result = select(*block('First Counsel LLP', 'same@example.org', 1),
                    *block('Other Counsel', 'same@example.org', 2))
    assert result['additional_emails'] == []


def test_report_appends_role_owner_pairs_and_metadata_survives(tmp_path, monkeypatch):
    result = select(*pair(), signature('Signing Counsel SA'))
    result.update(row=1, patent_id='WO/2025/123456', applicant='Source Applicant', priority_date='Source date')
    monkeypatch.setattr(agent, 'PCT_OUTPUT_DIR', tmp_path)
    path = agent.generate_work_report([result])
    book = openpyxl.load_workbook(path); sheet = book.active
    assert sheet.max_column == 13 and sheet.cell(1, 13).value == 'Additional Email'
    assert sheet.cell(2, 8).value == 'filing@signing.example'
    assert sheet.cell(2, 12).value == 'Source date'
    assert sheet.cell(2, 13).value == 'Agent email (First Counsel LLP): person@first.example'
    assert sheet.cell(2, 13).alignment.wrap_text and sheet.cell(2, 13).data_type == 's'
    assert ai.result_metadata(result)['additional_emails'] == result['additional_emails']
    book.close()


def test_signature_table_separates_firm_and_person_without_borrowing_contact_values():
    scan = {'pages': [{'page': 1, 'width': 600, 'height': 800}], 'lines': [
        line(1, 'Signature of applicant, agent or', 90, 300, width=130),
        line(2, 'Name', 90, 330, width=30), line(3, 'Signing Counsel SA', 245, 330),
        line(4, 'Name of signatory', 90, 350, width=90), line(5, 'Alex Green', 245, 350),
        line(6, 'Capacity', 90, 370), line(7, 'Agent', 245, 370),
        line(8, 'FOR RECEIVING OFFICE USE ONLY', 200, 410, width=200)]}
    found = signer.labelled_values(scan)
    assert [(f['kind'], f['text']) for f in found] == [('signer_entity', 'Signing Counsel SA'), ('signer_name', 'Alex Green')]
    scan['lines'][2]['rect'][1] += 10
    assert not any(item['kind'] == 'signer_entity' for item in signer.labelled_values(scan))


def test_signature_needs_both_independent_full_page_readings():
    first = ai.VisualExtraction(fields=[*pair(), signature('Signing Counsel SA')], reviewed_pages=[1, 2], uncertainties=[])
    second = first.model_copy(update={'fields': pair()})
    agreed, _ = policy.agree(first, second)
    assert ai.select_contacts(agreed, ai.load_rules(), [1, 2])['emails'] == ['person@first.example']
