import json
from pathlib import Path

import fitz
import pytest

from agents.pct_agent import ai_verifier as ai
from agents.pct_agent import pdf_extractor


@pytest.fixture(autouse=True)
def controlled_dns(monkeypatch):
    monkeypatch.setattr(ai.policy, 'email_domain_status', lambda email: 'mx')


def field(value="johnl@example.com", **changes):
    data = dict(kind="email", value=value, entity="John Lee", role="agent", page=1,
                evidence="Email: " + value, legible=True, block_id="page1-agent1",
                section="Agent / representative", entity_type="law_firm")
    data.update(changes)
    return ai.ContactField(**data)


def extraction(*fields, uncertainties=None):
    return ai.VisualExtraction(fields=list(fields), reviewed_pages=[1], uncertainties=uncertainties or [])


def pdf(tmp_path):
    path = tmp_path / "filing.pdf"
    with fitz.open() as doc:
        page = doc.new_page()
        page.insert_text((50, 60), "Agent: John Lee\nEmail: johnl@example.com\nTelephone: +1 212 555 0181", fontsize=16)
        doc.save(path)
    return path


def test_image_read_can_correct_plausible_but_wrong_ocr(tmp_path, monkeypatch):
    monkeypatch.setenv("PCT_AI_EVIDENCE_DIR", str(tmp_path / "evidence"))
    monkeypatch.setattr(ai, "_request", lambda pages, rules, context: extraction(field()))
    result = ai.verify_contacts(pdf(tmp_path), {"emails": ["john1@example.com"]})
    assert result["emails"] == ["johnl@example.com"]
    record = json.loads(Path(result["evidence_file"]).read_text())
    assert record["ocr_candidates"]["emails"] == ["john1@example.com"]
    assert Path(result["evidence_file"]).with_name("source.pdf").exists()


@pytest.mark.parametrize("candidate", [field(legible=False), field(page=5), field(evidence="Email unreadable"), field(entity="")])
def test_unsupported_fields_are_not_exported(candidate):
    result = ai.select_contacts(extraction(candidate), ai.load_rules(), [1])
    assert result["ai_status"] == "needs_review"
    assert result["emails"] == []
    assert result["status"] != "found"


def test_office_contacts_and_fax_are_excluded():
    result = ai.select_contacts(extraction(field("pct@wipo.int", role="patent_office"),
                                          field("+1 212 555 0100", kind="fax")), ai.load_rules(), [1])
    assert result["ai_status"] == "no_wanted_contacts"
    assert result["phones"] == []


def test_role_rules_apply_even_if_model_ignores_them():
    rules = ai.load_rules().model_copy(update={"include_roles": ["applicant"]})
    assert ai.select_contacts(extraction(field()), rules, [1])["emails"] == []


def test_different_owners_are_held_for_review():
    result = ai.select_contacts(extraction(field(), field("other@example.com", entity="Other Person")), ai.load_rules(), [1])
    assert result["ai_status"] == "needs_review"


def test_email_matching_is_case_insensitive_but_preserves_source_spelling():
    candidate = field("DocketDept@Example.COM", evidence="Email: docketdept@example.com")
    result = ai.select_contacts(extraction(candidate), ai.load_rules(), [1])
    assert result["emails"] == ["DocketDept@Example.COM"]


def test_same_contact_survives_firm_vs_attorney_owner_labels_between_reads():
    first = extraction(field("docket@example.com", entity="Example IP LLP"))
    second = extraction(field("docket@example.com", entity="Jane Attorney"))
    confirmed, issues = ai.policy.agree(first, second)
    assert [item.value for item in confirmed.fields] == ["docket@example.com"]
    assert "owner label differed" in "; ".join(issues)


def test_named_attorney_and_firm_fields_can_share_one_printed_block():
    result = ai.select_contacts(
        extraction(
            field("docket@example.com", entity="Example IP LLP"),
            field("Jane Attorney", kind="name", entity="Jane Attorney",
                  evidence="Name: Jane Attorney"),
        ),
        ai.load_rules(), [1],
    )
    assert result["emails"] == ["docket@example.com"]


def test_explicit_email_list_is_split_without_guessing():
    candidate = field("one@example.com; two@example.com",
                      evidence="Email: one@example.com; two@example.com")
    result = ai.select_contacts(extraction(candidate), ai.load_rules(), [1])
    assert result["emails"] == ["one@example.com", "two@example.com"]


@pytest.mark.parametrize("value,evidence", [
    ("H.Coleman <cosud@erols.com>", "Email: H.Coleman <cosud@erols.com>"),
    ("H.Coleman@cosud@erols.com", "Email: H.Coleman@cosud@erols.com"),
])
def test_display_name_email_syntax_keeps_only_bracketed_address(value, evidence):
    result = ai.select_contacts(extraction(field(value, evidence=evidence)), ai.load_rules(), [1])
    assert result["emails"] == ["cosud@erols.com"]


def test_unique_applicant_wins_even_when_multiple_agent_blocks_differ():
    result = ai.select_contacts(
        extraction(
            field("applicant@example.com", role="applicant", entity="Applicant Co",
                  block_id="applicant1", section="II Applicant", entity_type="company"),
            field("agent1@example.com", block_id="agent1"),
            field("agent2@example.com", entity="Other Agent", block_id="agent2"),
        ),
        ai.load_rules(), [1],
    )
    assert result["emails"] == ["applicant@example.com"]
    assert result["contact_role"] == "applicant"


def test_primary_iv_1_agent_wins_over_additional_iv_2_agent():
    result = ai.select_contacts(
        extraction(
            field("primary@example.com", block_id="agent1", section="IV-1 Agent"),
            field("additional@example.com", entity="Additional Agent", block_id="agent2",
                  section="IV-2 Additional agent(s)"),
        ),
        ai.load_rules(), [1],
    )
    assert result["emails"] == ["primary@example.com"]


def test_earliest_applicant_block_is_primary_when_continuation_applicant_differs():
    source = ai.VisualExtraction(
        fields=[
            field("primary@example.com", role="applicant", entity="Primary Applicant",
                  block_id="applicant1", section="II Applicant", page=1,
                  entity_type="company"),
            field("further@example.com", role="applicant", entity="Further Applicant",
                  block_id="applicant2", section="II Further Applicant", page=2,
                  entity_type="company"),
        ],
        reviewed_pages=[1, 2], uncertainties=[],
    )
    result = ai.select_contacts(source, ai.load_rules(), [1, 2])
    assert result["emails"] == ["primary@example.com"]


def test_agent_email_is_used_when_applicant_contact_has_no_email():
    result = ai.select_contacts(
        extraction(
            field("+1 212 555 0100", kind="phone", role="applicant", entity="Applicant Co",
                  block_id="applicant1", section="II Applicant", entity_type="company",
                  evidence="Telephone: +1 212 555 0100"),
            field("agent@example.com", block_id="agent1"),
        ),
        ai.load_rules(), [1],
    )
    assert result["emails"] == ["agent@example.com"]
    assert result["contact_role"] == "agent"


def test_shared_email_is_retained_when_address_cannot_decide_owner():
    result = ai.select_contacts(
        extraction(
            field("shared@example.com", role="applicant", entity="Applicant Co",
                  block_id="applicant1", section="II Applicant", entity_type="company"),
            field("shared@example.com", block_id="agent1"),
        ),
        ai.load_rules(), [1],
    )
    assert result["emails"] == ["shared@example.com"]
    assert result["contact_role"] == "shared"
    assert result["category"] == ""


def test_one_email_survives_ambiguous_agent_names_or_phones():
    result = ai.select_contacts(
        extraction(
            field("shared@example.com", entity="First Attorney", block_id="agent1"),
            field("shared@example.com", entity="Second Attorney", block_id="agent2"),
            field("+1 212 555 0181", kind="phone", entity="Second Attorney",
                  block_id="agent2", evidence="Telephone: +1 212 555 0181"),
        ),
        ai.load_rules(), [1],
    )
    assert result["emails"] == ["shared@example.com"]
    assert result["phones"] == ["+1 212 555 0181"]


def test_one_image_read_plus_exact_ocr_can_confirm_email():
    first = extraction(field("docket@example.com"))
    second = extraction(field("dockel@example.com"))
    confirmed, _ = ai.policy.agree(first, second, ["docket@example.com"])
    assert [item.value for item in confirmed.fields] == ["docket@example.com"]


def test_focused_third_read_resolves_one_character_ocr_conflict(tmp_path, monkeypatch):
    monkeypatch.setenv("PCT_AI_EVIDENCE_DIR", str(tmp_path / "evidence"))
    readings = iter([
        extraction(field("sean@boostreadmills.com")),
        extraction(field("sean@boostreadmills.com")),
        extraction(field("sean@boosttreadmills.com")),
    ])
    monkeypatch.setattr(ai, "_request", lambda pages, rules, context: next(readings))
    result = ai.verify_contacts(pdf(tmp_path), {"emails": ["sean@boosttreadmills.com"]})
    assert result["emails"] == ["sean@boosttreadmills.com"]
    assert "third image reading confirmed OCR" in result["reason"]


def test_focused_confirmation_keeps_original_contact_block_identity():
    original_email = field("aslewis@jonesday.com", block_id="agent2")
    confirmed = extraction(
        original_email,
        field("+1 212 326 3939", kind="phone", block_id="agent2",
              evidence="Telephone: +1 212 326 3939"),
        field("Applicant Co", kind="name", role="applicant", entity="Applicant Co",
              block_id="agent1", section="II Applicant", entity_type="company",
              evidence="Name: Applicant Co"),
    )
    focused = extraction(field("aslewis@jonesday.com", block_id="agent1"))
    conflict = [(ai.policy.agreement_key(original_email), ["asEewis@jonesday.com"])]
    resolved, _ = ai._resolve_ocr_conflicts(confirmed, focused, conflict)
    result = ai.select_contacts(resolved, ai.load_rules(), [1])
    assert result["emails"] == ["aslewis@jonesday.com"]
    assert result["phones"] == ["+1 212 326 3939"]


def test_focused_read_can_match_one_of_two_disputed_spellings(tmp_path, monkeypatch):
    monkeypatch.setenv("PCT_AI_EVIDENCE_DIR", str(tmp_path / "evidence"))
    readings = iter([
        extraction(field("farmini@shackelford.law")),
        extraction(field("famini@shackelford.law")),
        extraction(field("famini@shackelford.law")),
    ])
    monkeypatch.setattr(ai, "_request", lambda pages, rules, context: next(readings))
    result = ai.verify_contacts(pdf(tmp_path), {"emails": []})
    assert result["emails"] == ["famini@shackelford.law"]
    assert "matched an earlier email spelling" in result["reason"]


def test_unclear_read_can_trigger_crop_but_cannot_export_without_confirmation(tmp_path, monkeypatch):
    monkeypatch.setenv("PCT_AI_EVIDENCE_DIR", str(tmp_path / "evidence"))
    readings = iter([
        extraction(field("farmini@shackelford.law")),
        extraction(field("famini@shackelford.law", legible=False)),
        extraction(field("famini@shackelford.law")),
    ])
    monkeypatch.setattr(ai, "_request", lambda pages, rules, context: next(readings))
    result = ai.verify_contacts(pdf(tmp_path), {"emails": []})
    assert result["emails"] == ["famini@shackelford.law"]


def test_new_focused_spelling_requires_a_matching_fourth_read(tmp_path, monkeypatch):
    monkeypatch.setenv("PCT_AI_EVIDENCE_DIR", str(tmp_path / "evidence"))
    readings = iter([
        extraction(field("nflix-docketing@gtlaw.com", role="applicant",
                         section="II Applicant", entity="Netflix", block_id="applicant1")),
        extraction(field("nfix-docketing@gtlaw.com", role="applicant",
                         section="II Applicant", entity="Netflix", block_id="applicant1")),
        extraction(field("nflx-docketing@gtlaw.com", role="applicant",
                         section="II Applicant", entity="Netflix", block_id="applicant1")),
        extraction(field("nflx-docketing@gtlaw.com", role="applicant",
                         section="II Applicant", entity="Netflix", block_id="applicant1")),
    ])
    monkeypatch.setattr(ai, "_request", lambda pages, rules, context: next(readings))
    result = ai.verify_contacts(pdf(tmp_path), {"emails": []})
    assert result["emails"] == ["nflx-docketing@gtlaw.com"]
    assert "Two focused image readings" in result["reason"]


def test_repeated_malformed_email_needs_two_valid_crop_reads(tmp_path, monkeypatch):
    monkeypatch.setenv("PCT_AI_EVIDENCE_DIR", str(tmp_path / "evidence"))
    malformed = "H.Coleman@cosud@erols.com"
    readings = iter([
        extraction(field(malformed)),
        extraction(field(malformed)),
        extraction(field("cosud@erols.com")),
        extraction(field("cosud@erols.com")),
    ])
    monkeypatch.setattr(ai, "_request", lambda pages, rules, context: next(readings))
    result = ai.verify_contacts(pdf(tmp_path), {"emails": []})
    assert result["emails"] == ["cosud@erols.com"]


def test_unresolved_applicant_email_blocks_agent_fallback(tmp_path, monkeypatch):
    monkeypatch.setenv("PCT_AI_EVIDENCE_DIR", str(tmp_path / "evidence"))
    agent = field("agent@example.com", block_id="agent1")
    readings = iter([
        extraction(field("appllcant@example.com", role="applicant", entity="Applicant Co",
                         block_id="applicant1", section="II Applicant", entity_type="company"), agent),
        extraction(field("applicant@examp1e.com", role="applicant", entity="Applicant Co",
                         block_id="applicant1", section="II Applicant", entity_type="company"), agent),
    ])
    monkeypatch.setattr(ai, "_request", lambda pages, rules, context: next(readings))
    result = ai.verify_contacts(pdf(tmp_path), {"emails": ["agent@example.com"]})
    assert result["emails"] == []
    assert "applicant-priority" in result["reason"]


def test_ai_failure_never_falls_back_to_unverified_ocr(tmp_path, monkeypatch):
    monkeypatch.setenv("PCT_AI_EVIDENCE_DIR", str(tmp_path / "evidence"))
    def fail(*args):
        raise RuntimeError("openai_http_429")
    monkeypatch.setattr(ai, "_request", fail)
    result = ai.verify_contacts(pdf(tmp_path), {"emails": ["john1@example.com"]})
    assert result["status"] == "error" and result["emails"] == []
    assert Path(result["evidence_file"]).exists()


def test_changed_rules_invalidate_verified_cache(tmp_path, monkeypatch):
    monkeypatch.setenv("PCT_AI_EVIDENCE_DIR", str(tmp_path / "evidence"))
    calls = []
    def request(*args):
        calls.append(1)
        return extraction(field())
    monkeypatch.setattr(ai, "_request", request)
    source = pdf(tmp_path)
    ai.verify_contacts(source, {})
    ai.verify_contacts(source, {})
    assert len(calls) == 2
    rules = ai.load_rules().model_dump()
    rules["instructions"].append("Only direct contact addresses")
    rules_path = tmp_path / "rules.json"
    rules_path.write_text(json.dumps(rules))
    monkeypatch.setenv("PCT_AI_RULES_FILE", str(rules_path))
    ai.verify_contacts(source, {})
    assert len(calls) == 4


def test_extractor_routes_through_ai_and_passes_context(monkeypatch):
    monkeypatch.setenv("PCT_AI_VERIFY_ENABLED", "true")
    monkeypatch.setattr(pdf_extractor, "_extract_ocr_contacts", lambda *args: {"emails": ["bad@example.com"]})
    def verify(path, candidates, step, context):
        assert context == {"id": "WO1"}
        return {"status": "not_found", "emails": [], "ai_status": "needs_review"}
    monkeypatch.setattr(ai, "verify_contacts", verify)
    assert pdf_extractor.extract_contacts_from_pdf("unused.pdf", context={"id": "WO1"})["emails"] == []


def test_applicant_cache_cannot_bypass_document_verification(monkeypatch):
    from agents.pct_agent.pipeline import ContactCache
    monkeypatch.setenv("PCT_AI_VERIFY_ENABLED", "true")
    cache = ContactCache()
    cache.put("Same Applicant", {"status": "found", "emails": ["old@example.com"]})
    assert cache.get("Same Applicant") is None


def test_legacy_progress_does_not_skip_ai_verification(tmp_path, monkeypatch):
    from agents.pct_agent.pipeline import ProgressFile
    monkeypatch.setenv("PCT_AI_VERIFY_ENABLED", "true")
    path = tmp_path / "progress.jsonl"
    path.write_text(json.dumps({"patent_id": "WO1", "status": "found"}) + "\n")
    assert ProgressFile.load_completed(path) == set()


def test_report_matches_aneeq_columns_exactly(tmp_path, monkeypatch):
    import openpyxl
    from agents.pct_agent import agent
    monkeypatch.setattr(agent, "PCT_OUTPUT_DIR", tmp_path)
    expected = ["Publication No", "Title", "Application No.", "Applicant", "Url", "Cat",
                "Phone No.", "Email", "Agent Name", "Country", "Researcher", "Priorty Date"]
    for kind in ("worked", "not_found"):
        path = agent.generate_work_report([{"row": 1, "patent_id": "WO1", "ai_status": "needs_review",
                                           "reason": "Unclear l/1", "evidence_file": "source.pdf",
                                           "researcher": "Aneeq", "priority_date": "05 August to 10 August 2026"}], kind=kind)
        wb = openpyxl.load_workbook(path)
        assert wb.active.max_column == 12
        assert [c.value for c in wb.active[1]] == expected
        assert wb.active.cell(2, 11).value == "Aneeq"
        assert wb.active.cell(2, 12).value == "05 August to 10 August 2026"
        assert wb.active.cell(2, 8).value is None


def test_quota_failure_is_reported_before_a_scrape_starts(monkeypatch):
    from agents.pct_agent import agent
    def unavailable(**kwargs):
        assert kwargs["check_api"] is True
        raise ValueError("openai_http_429_credit_balance_exhausted")
    monkeypatch.setattr(agent, "validate_configuration", unavailable)
    result = agent.run_agent({"mode": "wipo_download"})
    assert result["status"] == "failure"
    assert "credit_balance_exhausted" in result["error"]


def test_machine_error_does_not_expose_raw_response():
    class Response:
        status_code = 429
        def json(self):
            return {"error": {"code": "credit_balance_exhausted", "message": "private request content"}}
    assert ai._api_error(Response()) == "openai_http_429_credit_balance_exhausted"
