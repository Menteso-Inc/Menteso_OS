"""Regression examples for the employee workflow: enrich one agent, preserve WIPO."""
import openpyxl
import pytest

from agents.pct_agent import agent, ai_verifier as ai
from agents.pct_agent.pipeline import _make_result


def contact(kind, value, **kwargs):
    data = {"kind": kind, "value": value, "entity": "Peter Counsel", "role": "agent",
            "page": 2, "evidence": kind + ": " + value, "legible": True,
            "block_id": "page2-agent1", "section": "IV. Agent / representative",
            "entity_type": "law_firm"}
    data.update(kwargs)
    return ai.ContactField(**data)


def select(*fields):
    return ai.select_contacts(ai.VisualExtraction(fields=list(fields), reviewed_pages=[1, 2],
                                                 uncertainties=[]), ai.load_rules(), [1, 2])


def agent_fields():
    return [contact("name", "Peter Counsel"), contact("email", "peterl01@example.org"),
            contact("phone", "+49 30 555 0101"), contact("country", "Germany")]


def test_applicant_and_inventor_details_never_fill_agent_columns():
    result = select(*agent_fields(),
                    contact("name", "Anne Applicant", role="applicant", entity="Anne Applicant",
                            section="II Applicant", block_id="applicant1"),
                    contact("email", "anne@example.com", role="applicant", entity="Anne Applicant",
                            section="II Applicant", block_id="applicant1"),
                    contact("phone", "+1 212 555 0199", role="inventor", block_id="inventor1"))
    assert result["agent_name"] == "Peter Counsel"
    assert result["emails"] == ["anne@example.com"]
    assert result["phones"] == []
    assert result["contact_role"] == "applicant"
    assert result["country"] == ""


def test_applicant_contact_preserved_without_relabeling_as_agent():
    result = select(contact("email", "applicant@example.com", role="applicant", section="II Applicant"))
    assert result["emails"] == ["applicant@example.com"] and result["agent_name"] == ""
    assert result["display_name"] == "Sir/Ma'am"
    assert result["contact_role"] == "applicant"


def test_mislabeled_applicant_section_is_rejected():
    result = select(contact("email", "applicant@example.com", section="II Applicant"))
    assert result["ai_status"] == "needs_review"
    assert result["emails"] == []


def test_name_cannot_be_borrowed_from_another_block():
    result = select(contact("email", "peter@example.org"),
                    contact("name", "Other Agent", entity="Other Agent", block_id="page2-agent2"))
    assert result["emails"] == ["peter@example.org"]
    assert result["agent_name"] == ""


def test_same_company_with_two_contact_blocks_keeps_unique_email_without_merging_phone():
    result = select(contact("email", "peter@example.org"),
                    contact("phone", "+49 30 555 0100", block_id="page3-agent2"))
    assert result["ai_status"] == "verified"
    assert result["emails"] == ["peter@example.org"] and result["phones"] == []
    assert result["contact_owner"] == ""


def test_different_name_inside_same_owner_block_is_not_exported():
    result = select(contact("email", "peter@example.org"), contact("name", "Anne Applicant"))
    assert result["emails"] == ["peter@example.org"]
    assert "mismatch" in result["reason"]
    assert result["agent_name"] == ""


def test_country_is_never_inferred_from_application_or_phone():
    result = select(contact("email", "peter@example.org"), contact("phone", "+49 30 555 0100"))
    assert result["country"] == ""
    assert ai._country_code("EP") == ""
    assert ai._country_code("Republic of Korea") == "KR"


def test_unknown_entity_type_does_not_guess_category():
    assert select(contact("email", "peter@example.org", entity_type="unknown"))["category"] == ""


@pytest.mark.parametrize("entity_type,category", [("law_firm", "Slf"), ("company", "Corp"), ("individual", "Ind")])
def test_category_spellings_match_reference(entity_type, category):
    assert select(contact("email", "peter@example.org", entity_type=entity_type))["category"] == category


def test_shuffled_wipo_columns_preserve_applicant(tmp_path):
    path = tmp_path / "wipo.xlsx"
    wb = openpyxl.Workbook()
    wb.active.append(["Applicant", "Appl.No", "Title", "ID", "IPC", "Kind"])
    wb.active.append(["EXAMPLE MANUFACTURING OÜ", "EP2024/123456", "Original title", "WO/2025/168220", "A01B", "A1"])
    wb.save(path)
    row = agent.read_input_excel(path)[0]
    assert row["applicant"] == "EXAMPLE MANUFACTURING OÜ"
    assert row["title"] == "Original title"
    assert row["appl_no"] == "EP2024/123456"


def test_missing_source_applicant_column_is_not_guessed():
    with pytest.raises(ValueError, match="Applicant"):
        agent._parse_source_rows([["ID", "Title", "Appl.No", "Other"],
                                  ["WO/2025/123456", "Title", "EP2024/123456", "Unrelated Name"]])


@pytest.mark.parametrize("builder", [agent._row_result, _make_result])
def test_enrichment_and_export_preserve_source_and_contact_ownership(tmp_path, monkeypatch, builder):
    monkeypatch.setenv("PCT_AI_VERIFY_ENABLED", "true")
    monkeypatch.setattr(agent, "PCT_OUTPUT_DIR", tmp_path)
    row = {"id": "WO/2025/168220", "title": "Original title", "appl_no": "EP2024/123456",
           "applicant": "EXAMPLE MANUFACTURING OÜ", "researcher": "Aneeq",
           "priority_date": "05 August to 10 August 2026", "_row_idx": 1}
    verified = select(*agent_fields())
    args = (row, "https://patentscope.wipo.int/source", "EP", "found")
    result = builder(1, *args) if builder is agent._row_result else builder(*args)
    assert result["country"] == ""
    result.update(emails=verified["emails"], phones=verified["phones"], name=verified["name"])
    result.update(ai.result_metadata(verified))
    path, _ = agent.write_pct_reports([result])
    wb = openpyxl.load_workbook(path)
    assert wb.active.max_column == 12
    assert [c.value for c in wb.active[2]] == [
        row["id"], row["title"], row["appl_no"], row["applicant"], args[1], "Slf",
        "+49 30 555 0101", "peterl01@example.org", "Peter Counsel", "DE", "Aneeq", row["priority_date"]]


def test_aneeq_input_layout_preserves_employee_metadata():
    row = ["WO/2025/168220", "Title", "EP2024/123456", "Company", "url", "Slf",
           "", "", "", "", "Aneeq", "05 August to 10 August 2026"]
    parsed = agent._parse_source_rows([agent.WORK_REPORT_HEADERS, row])[0]
    assert parsed["applicant"] == "Company"
    assert parsed["researcher"] == "Aneeq"
    assert parsed["priority_date"] == row[-1]
