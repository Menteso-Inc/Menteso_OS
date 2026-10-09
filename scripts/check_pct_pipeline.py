"""Offline evidence replay and timing simulation; no WIPO or OpenAI calls.

Run with the existing venv from this checkout. Supply a baseline ai_verifier.py
and local evidence directory. The report contains counts, never contact data.
"""
import argparse
import importlib.util
import json
import sys
import tempfile
import time
from copy import copy
from pathlib import Path
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from agents.pct_agent import ai_verifier as candidate, agent, overlap


def load_baseline(path, name="_baseline_verifier"):
    spec = importlib.util.spec_from_file_location("agents.pct_agent." + name, path)
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def replay(module, record, source_pdf, output):
    readings = [record["visual_extraction"], record["second_reading"]]
    focused = [record[key] for key in ("tie_break_reading", "fourth_reading") if record.get(key)]
    # Initial read identities are explicit; separate barrier tests exercise the
    # real concurrent execution and prove completion order cannot swap them.
    request_records = iter(focused if hasattr(module, "_read_original_pages") else readings + focused)
    def request(*args):
        return module.VisualExtraction.model_validate(next(request_records))
    def read_pair(*args):
        return tuple(module.VisualExtraction.model_validate(item) for item in readings)
    saved_dns = record["result"].get("email_checks", {})
    page_numbers = record["pages_reviewed"]
    from contextlib import ExitStack
    with ExitStack() as stack:
        for name, replacement in {
            "load_rules": lambda: module.ContactRules.model_validate(record["rules"]),
            "_request": request,
            "_render_pages": lambda _: ([(number, b"original-image") for number in page_numbers], record["total_pages"]),
            "_render_focus_crops": lambda _, numbers: [(number, b"focused-image") for number in numbers],
            "_high_res_ocr_emails": lambda *_: record.get("high_resolution_ocr_emails", []),
        }.items():
            stack.enter_context(patch.object(module, name, replacement))
        if hasattr(module, "_read_original_pages"):
            stack.enter_context(patch.object(module, "_read_original_pages", read_pair))
        stack.enter_context(patch.object(module.policy, "email_domain_status",
                                         lambda email: saved_dns.get(email, {}).get("domain", "unconfirmed")))
        stack.enter_context(patch.dict("os.environ", {
            "PCT_AI_EVIDENCE_DIR": str(output), "PCT_EMAIL_DNS_CHECK": "true",
            "PCT_AI_VERIFY_ENABLED": "true",
        }))
        result = module.verify_contacts(source_pdf, record.get("ocr_candidates", {}))
    return {key: value for key, value in result.items() if key not in {"evidence_file", "verification_revision"}}


def compare_evidence(baseline, baseline_agent, evidence_dir, limit):
    eligible = []
    for path in evidence_dir.glob("*/verification.json"):
        # Do not inspect an evidence record currently being written by a run.
        if time.time() - path.stat().st_mtime < 60:
            continue
        try:
            record = json.loads(path.read_text(encoding="utf-8"))
            if all(record.get(key) for key in ("visual_extraction", "second_reading", "result")) and path.with_name("source.pdf").exists():
                eligible.append((path, record))
        except (OSError, ValueError):
            continue
    eligible.sort(key=lambda pair: pair[0].stat().st_mtime, reverse=True)
    chosen = eligible[:limit]
    if not chosen:
        raise ValueError("No complete evidence records found")
    mismatches = []
    old_rows, new_rows = [], []
    workbook_cells_match = True
    with tempfile.TemporaryDirectory(prefix="pct-evidence-replay-") as scratch:
        scratch = Path(scratch)
        for index, (path, record) in enumerate(chosen, 1):
            previous = replay(baseline, record, path.with_name("source.pdf"), scratch / f"old-{index}")
            updated = replay(candidate, record, path.with_name("source.pdf"), scratch / f"new-{index}")
            if previous != updated:
                mismatches.append({"record": path.parent.name,
                                   "baseline_error": previous.get("error"),
                                   "candidate_error": updated.get("error"),
                                   "fields": [key for key in set(previous) | set(updated) if previous.get(key) != updated.get(key)]})
            # Synthetic source metadata keeps customer identifiers out of the report.
            row = {"row": index, "patent_id": f"WO-REPLAY-{index}", "title": f"Source row {index}",
                   "appl_no": f"APP-{index}", "applicant": f"Source applicant {index}",
                   "researcher": "Original researcher", "priority_date": "Original date"}
            old_rows.append({**row, **previous})
            new_rows.append({**row, **updated})
        import openpyxl
        original_output = agent.PCT_OUTPUT_DIR
        original_baseline_output = baseline_agent.PCT_OUTPUT_DIR
        try:
            baseline_agent.PCT_OUTPUT_DIR = scratch / "old-sheets"
            baseline_agent.PCT_OUTPUT_DIR.mkdir(parents=True)
            old_paths = baseline_agent.write_pct_reports(old_rows)
            agent.PCT_OUTPUT_DIR = scratch / "new-sheets"
            new_paths = agent.write_pct_reports(new_rows)
            for old_path, new_path in zip(old_paths, new_paths):
                old_book, new_book = openpyxl.load_workbook(old_path), openpyxl.load_workbook(new_path)
                try:
                    workbook_cells_match = workbook_cells_match and list(old_book.active.values) == list(new_book.active.values)
                    assert old_book.active.max_column == new_book.active.max_column == 12
                    assert old_book.active.title == new_book.active.title
                    assert [(copy(cell.font), cell.number_format, copy(cell.alignment)) for row in old_book.active for cell in row] == [(copy(cell.font), cell.number_format, copy(cell.alignment)) for row in new_book.active for cell in row]
                    assert {k: v.width for k, v in old_book.active.column_dimensions.items()} == {k: v.width for k, v in new_book.active.column_dimensions.items()}
                finally:
                    old_book.close()
                    new_book.close()
        finally:
            agent.PCT_OUTPUT_DIR = original_output
            baseline_agent.PCT_OUTPUT_DIR = original_baseline_output
    return {"records_compared": len(chosen), "identical_results": len(chosen) - len(mismatches),
            "mismatches": mismatches, "workbook_cells_match": workbook_cells_match,
            "scope": "Same recorded AI/OCR/DNS evidence, full verification and selection replay; no fresh API accuracy claim"}


def timing_simulation(count=20, scale=0.01):
    # Measured 17-row averages: download 6.4s, local work 1.5s, two AI reads 19.9s.
    def request(*args):
        time.sleep(9.95 * scale)
        return None
    rows = [{"id": f"WO/{i}", "title": "Timing fixture", "appl_no": "US1", "applicant": "Fixture"}
            for i in range(count)]
    class Browser:
        def scrape_patent(self, *args, **kwargs):
            time.sleep(6.4 * scale)
            return "simulated.pdf"
    def verify(*args, **kwargs):
        time.sleep(1.5 * scale)
        candidate._read_original_pages([(1, b"fixture")], None, {})
        return {"status": "found", "emails": ["fixture@example.org"]}
    with patch.object(candidate, "_request", request):
        started = time.perf_counter()
        for _ in rows:
            Browser().scrape_patent()
            time.sleep(1.5 * scale)
            request()
            request()
        serial = time.perf_counter() - started
        started = time.perf_counter()
        result = overlap.process_rows(
            rows, Browser(), verify,
            lambda row, url, contacts: {"row": row["_row_idx"], **contacts},
            lambda _: None, lambda _: None, lambda _: None, lambda: False,
            "simulation", max_pending=2, min_interval=10 * scale,
        )
        overlapped = time.perf_counter() - started
    assert len(result) == count
    return {"rows": count, "sequential_wall_seconds": round(serial, 3),
            "overlapped_wall_seconds": round(overlapped, 3), "speedup": round(serial / overlapped, 2),
            "scope": "Scaled delay simulation; actual WIPO/OpenAI performance is unmeasured"}


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--baseline-verifier", type=Path, required=True)
    parser.add_argument("--evidence-dir", type=Path, required=True)
    parser.add_argument("--limit", type=int, default=100)
    parser.add_argument("--report", type=Path, required=True)
    args = parser.parse_args()
    baseline = load_baseline(args.baseline_verifier)
    baseline_agent = load_baseline(args.baseline_verifier.with_name("agent.py"), "_baseline_agent")
    # Any accidental HTTP call is a test failure, protecting the active run.
    with patch("requests.post", side_effect=AssertionError("Network disabled during offline validation")):
        report = {"evidence": compare_evidence(baseline, baseline_agent, args.evidence_dir, args.limit),
                  "timing": timing_simulation()}
    args.report.parent.mkdir(parents=True, exist_ok=True)
    args.report.write_text(json.dumps(report, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, indent=2))
    if report["evidence"]["mismatches"] or not report["evidence"]["workbook_cells_match"]:
        raise SystemExit(1)
